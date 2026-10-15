"""Repository dokumen TSD lokal: preview, download, dan lifecycle sumber."""

from __future__ import annotations

import json
import hashlib
import shutil
import subprocess
import tempfile
import threading
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from docx import Document

from app.indexer import build_index
from app.indexer.ingest import PARSED_DIR, SAMPLES_DIR, validate_docx, validate_filename
from app.parsers import docx_parser
from app.modules import infer_module
from app.banks import UNKNOWN_BANK, bank_options, detect_bank, normalize_bank
from app.source_paths import paths_for


BASE_DIR = Path(__file__).resolve().parents[2]
ACTIVE_DIR = BASE_DIR / SAMPLES_DIR
PARSED_ROOT = BASE_DIR / PARSED_DIR
INACTIVE_DIR = BASE_DIR / "storage" / "documents" / "inactive"
REGISTRY_PATH = BASE_DIR / "storage" / "documents" / "registry.json"
PDF_DIR = BASE_DIR / "storage" / "documents" / "pdf"
_lock = threading.Lock()
_pdf_lock = threading.Lock()


def _detected_module(item: dict[str, Any], bank: str = "SCB") -> str:
    """Deteksi modul dari identitas dokumen dan isi hasil parser."""
    filename = item.get("tsd_filename") or item.get("filename") or ""
    segment = item.get("segment") or ""
    identity_module = infer_module(segment, filename)
    if identity_module != "LAINNYA":
        return identity_module
    parsed_path = paths_for(bank).parsed_documents / f"{Path(filename).stem}.json"
    content: list[str] = []
    try:
        parsed = json.loads(parsed_path.read_text(encoding="utf-8"))
        for procedure in parsed.get("procedures", []):
            content.extend((
                procedure.get("sp_name") or "",
                (procedure.get("data_model") or {}).get("section") or "",
                (procedure.get("data_flow") or {}).get("section") or "",
            ))
    except (OSError, json.JSONDecodeError):
        pass
    return infer_module(*content)


def _detected_bank(filename: str, path: Path | None = None) -> str:
    text = ""
    source = path or (ACTIVE_DIR / filename)
    try:
        document = Document(source)
        text = "\n".join(
            " ".join(paragraph.text.split())
            for paragraph in document.paragraphs[:160]
            if paragraph.text.strip()
        )
    except (OSError, ValueError):
        pass
    return str(detect_bank(filename, text)["suggestion"])


def _scope_image_paths(parsed: dict[str, Any], bank: str) -> None:
    if bank == "SCB":
        return
    for procedure in parsed.get("procedures", []):
        for block_name in ("data_model", "data_flow"):
            for image in (procedure.get(block_name) or {}).get("images", []):
                if not image["path"].startswith(f"{bank}/"):
                    image["path"] = f"{bank}/{image['path']}"


def _read_registry(bank: str = "SCB") -> dict[str, Any]:
    registry_path = paths_for(bank).registry
    if not registry_path.exists():
        return {"documents": {}}
    try:
        return json.loads(registry_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return {"documents": {}}


def _write_registry(registry: dict[str, Any], bank: str = "SCB") -> None:
    registry_path = paths_for(bank).registry
    registry_path.parent.mkdir(parents=True, exist_ok=True)
    temp = registry_path.with_suffix(".tmp")
    temp.write_text(json.dumps(registry, ensure_ascii=False, indent=2), encoding="utf-8")
    temp.replace(registry_path)


def _resolve(filename: str, bank: str = "SCB", *, allow_inactive: bool = True) -> tuple[Path, str]:
    name = validate_filename(filename)
    paths = paths_for(bank).ensure()
    active = paths.documents / name
    if active.is_file():
        return active, "active"
    inactive = paths.inactive_documents / name
    if allow_inactive and inactive.is_file():
        return inactive, "inactive"
    raise FileNotFoundError("Dokumen TSD tidak ditemukan.")


def document_path(filename: str, bank: str = "SCB") -> tuple[Path, str]:
    path, status = _resolve(filename, bank)
    validate_docx(path)
    return path, status


class PdfConverterUnavailableError(RuntimeError):
    pass


def pdf_path(filename: str, bank: str = "SCB") -> Path:
    """Konversi DOCX ke PDF visual dan cache sampai file sumber berubah."""
    source, _ = document_path(filename, bank)
    pdf_dir = paths_for(bank).ensure().pdf
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        raise PdfConverterUnavailableError(
            "Preview visual memerlukan LibreOffice. Instal dengan: brew install --cask libreoffice"
        )
    pdf_dir.mkdir(parents=True, exist_ok=True)
    # Versi cache ikut berubah saat konfigurasi font/converter diperbarui.
    digest = hashlib.sha256(f"pdf-v2:{source.name}".encode()).hexdigest()[:16]
    destination = pdf_dir / f"{digest}.pdf"
    with _pdf_lock:
        if destination.exists() and destination.stat().st_mtime >= source.stat().st_mtime:
            return destination
        with tempfile.TemporaryDirectory(prefix="tsd-pdf-") as temp_dir:
            temp_root = Path(temp_dir)
            profile = temp_root / "libreoffice-profile"
            profile.mkdir()
            result = subprocess.run(
                [
                    executable,
                    f"-env:UserInstallation={profile.as_uri()}",
                    "--headless",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    temp_dir,
                    str(source),
                ],
                capture_output=True,
                text=True,
                timeout=180,
                check=False,
            )
            converted = temp_root / f"{source.stem}.pdf"
            if result.returncode != 0 or not converted.exists():
                detail = (result.stderr or result.stdout).strip()
                raise RuntimeError(f"Konversi PDF gagal. {detail}".strip())
            # /tmp dan volume storage adalah mount berbeda di Docker. Salin dulu ke
            # file sementara pada volume tujuan, lalu rename atomik di mount yang sama.
            with tempfile.NamedTemporaryFile(dir=pdf_dir, suffix=".pdf.tmp", delete=False) as output:
                staged = Path(output.name)
            try:
                shutil.copyfile(converted, staged)
                staged.replace(destination)
            finally:
                staged.unlink(missing_ok=True)
    return destination


def preview(filename: str, bank: str = "SCB", offset: int = 0, limit: int = 80) -> dict[str, Any]:
    path, status = document_path(filename, bank)
    document = Document(path)
    blocks: list[dict[str, Any]] = []
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue
        style = paragraph.style.name if paragraph.style else "Normal"
        heading = style.lower().replace(" ", "").startswith("heading")
        blocks.append({"type": "heading" if heading else "paragraph", "style": style, "text": text})
    for table in document.tables:
        rows = [[" ".join(cell.text.split()) for cell in row.cells] for row in table.rows]
        if rows:
            blocks.append({"type": "table", "rows": rows[:100]})
    props = document.core_properties
    registry = _read_registry(bank).get("documents", {}).get(path.name, {})
    segment = registry.get("segment") or path.stem
    total_blocks = len(blocks)
    return {
        "filename": path.name,
        "status": status,
        "segment": segment,
        "module": registry.get("module") or infer_module(segment),
        "bank": bank,
        "size_bytes": path.stat().st_size,
        "modified_at": path.stat().st_mtime,
        "title": props.title or None,
        "author": props.author or None,
        "blocks": blocks[offset:offset + limit],
        "total_blocks": total_blocks,
        "offset": offset,
        "limit": limit,
    }


def sync_registry(documents: list[dict[str, Any]]) -> None:
    with _lock:
        now = datetime.now(timezone.utc).isoformat(timespec="seconds")
        for bank in bank_options():
            registry = _read_registry(bank)
            records = registry.setdefault("documents", {})
            for item in (row for row in documents if row.get("bank") == bank):
                record = records.setdefault(item["tsd_filename"], {})
                detected = _detected_module(item, bank)
                current = record.get("module")
                record.update({
                    "status": "active",
                    "segment": item.get("segment"),
                    "module": detected if not current or current == "LAINNYA" else current,
                    "bank": bank,
                    "bank_source": "selected_on_upload",
                    "updated_at": record.get("updated_at") or now,
                    "procedure_count": item.get("procedure_count"),
                })
            _write_registry(registry, bank)


def enrich_documents(documents: list[dict[str, Any]]) -> list[dict[str, Any]]:
    enriched = []
    for item in documents:
        row = dict(item)
        bank = row.get("bank") or "SCB"
        registry = _read_registry(bank).get("documents", {})
        record = registry.get(row["tsd_filename"], {})
        row["module"] = record.get("module") or _detected_module(row, bank)
        row["bank"] = bank
        row["tags"] = record.get("tags", [])
        row["updated_at"] = record.get("updated_at")
        enriched.append(row)
    return enriched


def update_metadata(
    filename: str, new_filename: str, module: str, bank: str, tags: list[str]
) -> dict[str, Any]:
    name = validate_filename(filename)
    renamed = validate_filename(new_filename.strip())
    clean_bank = normalize_bank(bank)
    paths = paths_for(clean_bank).ensure()
    source, status = _resolve(name, clean_bank)
    clean_module = module.strip().upper()
    if not clean_module or len(clean_module) > 40:
        raise ValueError("Modul wajib diisi dan maksimal 40 karakter.")
    clean_tags = list(dict.fromkeys(tag.strip() for tag in tags if tag.strip()))
    if len(clean_tags) > 10 or any(len(tag) > 30 for tag in clean_tags):
        raise ValueError("Maksimal 10 tag, masing-masing 30 karakter.")
    with _lock:
        destination = source.with_name(renamed)
        if renamed != name and ((paths.documents / renamed).exists() or (paths.inactive_documents / renamed).exists()):
            raise FileExistsError("Dokumen dengan nama file tersebut sudah ada.")

        old_parsed = paths.parsed_documents / f"{Path(name).stem}.json"
        old_parsed_hold = PARSED_ROOT / f"{Path(name).stem}.json.inactive"
        old_parsed_source = old_parsed if status == "active" else old_parsed_hold
        new_parsed = paths.parsed_documents / f"{Path(renamed).stem}.json"
        new_parsed_hold = PARSED_ROOT / f"{Path(renamed).stem}.json.inactive"
        parsed_destination = new_parsed if status == "active" else new_parsed_hold
        parsed_backup = old_parsed_source.with_name(f".{old_parsed_source.name}.rename-backup")

        if renamed != name:
            parsed_backup.unlink(missing_ok=True)
            source.replace(destination)
            if old_parsed_source.exists():
                old_parsed_source.replace(parsed_backup)
            try:
                if status == "active":
                    parsed = docx_parser.parse(destination, paths.images)
                    _scope_image_paths(parsed, clean_bank)
                    new_parsed.write_text(
                        json.dumps(parsed, indent=2, ensure_ascii=False), encoding="utf-8"
                    )
                    build_index.build()
                elif parsed_backup.exists():
                    parsed_backup.replace(parsed_destination)
            except BaseException:
                destination.replace(source)
                new_parsed.unlink(missing_ok=True)
                new_parsed_hold.unlink(missing_ok=True)
                if parsed_backup.exists():
                    parsed_backup.replace(old_parsed_source)
                raise
            parsed_backup.unlink(missing_ok=True)

        registry = _read_registry(clean_bank)
        records = registry.setdefault("documents", {})
        record = records.pop(name, {}) if renamed != name else records.setdefault(name, {})
        record.update({
            "module": clean_module,
            "bank": clean_bank,
            "bank_source": "manual",
            "tags": clean_tags,
            "status": status,
            "updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        })
        records[renamed] = record
        _write_registry(registry, clean_bank)
        if status == "active":
            build_index.build()
    return {"filename": renamed, "module": clean_module, "bank": clean_bank, "tags": clean_tags}


def set_bank(filename: str, bank: str, rebuild: bool = True) -> str:
    name = validate_filename(filename)
    clean_bank = normalize_bank(bank)
    with _lock:
        registry = _read_registry(clean_bank)
        record = registry.setdefault("documents", {}).setdefault(name, {})
        record.update({"bank": clean_bank, "bank_source": "manual", "updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})
        _write_registry(registry, clean_bank)
        if rebuild:
            build_index.build()
    return clean_bank


def list_banks() -> list[str]:
    return bank_options()


def inactive_documents(bank: str = "SCB") -> list[dict[str, Any]]:
    paths = paths_for(bank).ensure()
    registry = _read_registry(bank).get("documents", {})
    rows = []
    for path in sorted(paths.inactive_documents.glob("*.docx")):
        record = registry.get(path.name, {})
        segment = record.get("segment") or path.stem
        detected = _detected_module({"tsd_filename": path.name, "segment": segment}, bank)
        module = record.get("module")
        if not module or module == "LAINNYA":
            module = detected
        rows.append({
            "tsd_filename": path.name,
            "segment": segment,
            "module": module,
            "bank": bank,
            "status": "inactive",
            "procedure_count": record.get("procedure_count"),
            "indexed_sp": 0,
            "size_bytes": path.stat().st_size,
            "updated_at": record.get("updated_at"),
            "tags": record.get("tags", []),
        })
    return rows


def set_inactive(filename: str, bank: str = "SCB") -> dict[str, Any]:
    name = validate_filename(filename)
    paths = paths_for(bank).ensure()
    source, status = _resolve(name, bank, allow_inactive=False)
    if status != "active":
        raise ValueError("Dokumen sudah nonaktif.")
    destination = paths.inactive_documents / name
    if destination.exists():
        raise FileExistsError("Versi nonaktif dengan nama yang sama sudah ada.")
    parsed = paths.parsed_documents / f"{Path(name).stem}.json"
    parsed_hold = parsed.with_suffix(".json.inactive")
    with _lock:
        shutil.move(str(source), destination)
        if parsed.exists():
            parsed.replace(parsed_hold)
        try:
            build_index.build()
        except BaseException:
            shutil.move(str(destination), source)
            if parsed_hold.exists():
                parsed_hold.replace(parsed)
            raise
        registry = _read_registry(bank)
        record = registry.setdefault("documents", {}).setdefault(name, {})
        record.update({"status": "inactive", "updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})
        _write_registry(registry, bank)
    return {"filename": name, "status": "inactive"}


def reactivate(filename: str, bank: str = "SCB") -> dict[str, Any]:
    name = validate_filename(filename)
    paths = paths_for(bank).ensure()
    source, status = _resolve(name, bank)
    if status != "inactive":
        raise ValueError("Dokumen sudah aktif.")
    destination = paths.documents / name
    if destination.exists():
        raise FileExistsError("Dokumen aktif dengan nama yang sama sudah ada.")
    parsed_hold = paths.parsed_documents / f"{Path(name).stem}.json.inactive"
    parsed = paths.parsed_documents / f"{Path(name).stem}.json"
    with _lock:
        shutil.move(str(source), destination)
        if parsed_hold.exists():
            parsed_hold.replace(parsed)
        try:
            build_index.build()
        except BaseException:
            shutil.move(str(destination), source)
            if parsed.exists():
                parsed.replace(parsed_hold)
            raise
        registry = _read_registry(bank)
        record = registry.setdefault("documents", {}).setdefault(name, {})
        record.update({"status": "active", "updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds")})
        _write_registry(registry, bank)
    return {"filename": name, "status": "active"}


def delete(filename: str, bank: str = "SCB") -> dict[str, Any]:
    name = validate_filename(filename)
    paths = paths_for(bank).ensure()
    source, status = _resolve(name, bank)
    trash = source.with_name(f".{source.name}.deleting")
    parsed = paths.parsed_documents / f"{Path(name).stem}.json"
    parsed_hold = paths.parsed_documents / f"{Path(name).stem}.json.inactive"
    parsed_source = parsed if parsed.exists() else parsed_hold
    parsed_trash = parsed_source.with_name(f".{parsed_source.name}.deleting")
    with _lock:
        source.replace(trash)
        if parsed_source.exists():
            parsed_source.replace(parsed_trash)
        try:
            if status == "active":
                build_index.build()
        except BaseException:
            trash.replace(source)
            if parsed_trash.exists():
                parsed_trash.replace(parsed_source)
            raise
        trash.unlink(missing_ok=True)
        parsed_trash.unlink(missing_ok=True)
        registry = _read_registry(bank)
        registry.setdefault("documents", {}).pop(name, None)
        _write_registry(registry, bank)
    return {"filename": name, "status": "deleted"}


def replace_from_editor(filename: str, temp_path: Path, bank: str = "SCB") -> dict[str, Any]:
    """Ganti DOCX aktif dari editor, lalu parse dan bangun ulang indeks."""
    name = validate_filename(filename)
    paths = paths_for(bank).ensure()
    source, status = _resolve(name, bank, allow_inactive=False)
    if status != "active":
        raise ValueError("Hanya dokumen aktif yang dapat diedit.")
    validate_docx(temp_path)
    backup = source.with_name(f".{source.name}.editor-backup")
    parsed_path = paths.parsed_documents / f"{source.stem}.json"
    parsed_backup = parsed_path.with_name(f".{parsed_path.name}.editor-backup")
    with _lock:
        backup.unlink(missing_ok=True)
        parsed_backup.unlink(missing_ok=True)
        source.replace(backup)
        if parsed_path.exists():
            parsed_path.replace(parsed_backup)
        try:
            temp_path.replace(source)
            parsed = docx_parser.parse(source, paths.images)
            _scope_image_paths(parsed, bank)
            parsed_path.write_text(
                json.dumps(parsed, indent=2, ensure_ascii=False), encoding="utf-8"
            )
            build_index.build()
        except BaseException:
            source.unlink(missing_ok=True)
            parsed_path.unlink(missing_ok=True)
            backup.replace(source)
            if parsed_backup.exists():
                parsed_backup.replace(parsed_path)
            raise
        backup.unlink(missing_ok=True)
        parsed_backup.unlink(missing_ok=True)
        paths.pdf.mkdir(parents=True, exist_ok=True)
        for cached_pdf in paths.pdf.glob("*.pdf"):
            cached_pdf.unlink(missing_ok=True)
        registry = _read_registry(bank)
        record = registry.setdefault("documents", {}).setdefault(name, {})
        record.update({
            "status": "active",
            "updated_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
            "procedure_count": parsed.get("procedure_count"),
        })
        _write_registry(registry, bank)
    return {"filename": name, "status": "active", "reindexed": True}
