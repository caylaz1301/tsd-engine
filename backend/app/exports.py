"""Ekspor operasional untuk matriks cakupan dan laporan analisis SP."""

from __future__ import annotations

import hashlib
import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.db import index_db
from app.modules import infer_module


BASE_DIR = Path(__file__).resolve().parents[1]
EXPORT_DIR = BASE_DIR / "storage" / "exports"


def matrix_workbook() -> Path:
    stats = index_db.stats()
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    destination = EXPORT_DIR / "Matriks_Cakupan_TSD.xlsx"
    wb = Workbook()
    summary = wb.active
    summary.title = "Cakupan per Modul"
    summary.append(["Modul", "SP Utama", "Cocok", "Tanpa TSD", "Cakupan (%)"])
    for item in stats["modules"]:
        summary.append([
            item["module"], item["total"], item["matched"], item["sql_only"],
            item["matched"] / item["total"] if item["total"] else 0,
        ])
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    summary.column_dimensions["A"].width = 22
    for col in "BCDE":
        summary.column_dimensions[col].width = 16
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in summary["E"][1:]:
        cell.number_format = "0.0%"

    missing = wb.create_sheet("SP tanpa TSD")
    missing.append(["Modul", "Stored Procedure", "Database", "Berkas SQL"])
    con = index_db.connect()
    try:
        rows = con.execute(
            "SELECT sp_name, sql_database, sql_file, segment FROM sp_index "
            "WHERE status = 'sql_only' AND variant IS NULL ORDER BY sql_database, sp_name"
        ).fetchall()
    finally:
        con.close()
    for row in rows:
        missing.append([
            infer_module(row["segment"], row["sql_database"], row["sp_name"]),
            row["sp_name"], row["sql_database"], row["sql_file"],
        ])
    missing.freeze_panes = "A2"
    missing.auto_filter.ref = missing.dimensions
    for width, col in zip((18, 48, 22, 44), "ABCD"):
        missing.column_dimensions[col].width = width
    for cell in missing[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
    wb.save(destination)
    return destination


def sp_report_pdf(sp: dict[str, Any], analysis: dict[str, Any]) -> Path:
    executable = shutil.which("soffice") or shutil.which("libreoffice")
    if executable is None:
        raise RuntimeError("LibreOffice tidak tersedia untuk membuat laporan PDF.")
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)
    digest = hashlib.sha256(f"{sp['sp_name']}:{analysis}".encode()).hexdigest()[:16]
    destination = EXPORT_DIR / f"Analisis_{sp['sp_name']}_{digest}.pdf"
    if destination.exists():
        return destination

    with tempfile.TemporaryDirectory(prefix="tsd-report-") as temp_dir:
        temp = Path(temp_dir)
        doc = Document()
        doc.add_heading(f"Laporan Analisis Stored Procedure", 0)
        doc.add_heading(sp["sp_name"], 1)
        doc.add_paragraph(f"Status: {sp['status']} | Database: {sp.get('sql_database') or '-'}")
        doc.add_heading("Ringkasan AI", 1)
        doc.add_paragraph(analysis.get("summary") or "Ringkasan tidak tersedia.")
        doc.add_heading("Tujuan", 2)
        doc.add_paragraph(analysis.get("purpose") or "Belum teridentifikasi.")
        doc.add_heading("Alur yang teridentifikasi", 2)
        for step in analysis.get("process_steps", []):
            doc.add_paragraph(step, style="List Number")
        doc.add_heading("Data dibaca", 2)
        for table in analysis.get("data_reads", []):
            doc.add_paragraph(table, style="List Bullet")
        doc.add_heading("Data ditulis", 2)
        for table in analysis.get("data_writes", []):
            doc.add_paragraph(table, style="List Bullet")
        doc.add_heading("Perlu diverifikasi", 2)
        notes = analysis.get("review_notes", []) or ["Tidak ada catatan tambahan."]
        for note in notes:
            doc.add_paragraph(note, style="List Bullet")
        for image in sp.get("images", [])[:2]:
            path = index_db.IMAGES_DIR / image["path"]
            if path.exists():
                doc.add_heading(image["kind"].replace("_", " ").title(), 2)
                doc.add_picture(str(path), width=Inches(6.3))
        source = temp / "report.docx"
        doc.save(source)
        result = subprocess.run(
            [executable, "--headless", "--convert-to", "pdf", "--outdir", str(temp), str(source)],
            capture_output=True, text=True, timeout=180, check=False,
        )
        converted = temp / "report.pdf"
        if result.returncode != 0 or not converted.exists():
            raise RuntimeError((result.stderr or result.stdout or "Konversi laporan gagal.").strip())
        shutil.copyfile(converted, destination)
    return destination
