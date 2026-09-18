"""Quality gate deterministik untuk dokumen TSD berformat DOCX."""

from __future__ import annotations

import json
import os
import re
import shutil
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document

from app import config as C
from app.indexer.ingest import validate_docx, validate_filename
from app.parsers import docx_parser


STAGING_DIR = Path("storage/quality/staging")
REPORT_DIR = Path("storage/quality/reports")
PREVIEW_IMAGES_DIR = Path("storage/quality/images")
DATE_FORMATS = ("%d %B %Y", "%d %b %Y", "%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y")
NUMBER_RE = re.compile(r"^\s*(\d+(?:\.[0-9A-Za-z]+)*\.?)\s+")
PLACEHOLDER_RE = re.compile(r"\b(?:lorem ipsum|todo|tbd|xxx|insert (?:text|diagram)|placeholder)\b", re.I)
CHECKER_VERSION = 4


def _result(code: str, label: str, passed: bool, summary: str, findings: list[str] | None = None) -> dict:
    return {"code": code, "label": label, "status": "pass" if passed else "fail", "summary": summary, "findings": findings or []}


def _parse_date(raw: str) -> datetime | None:
    cleaned = " ".join(raw.replace(",", " ").split())
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(cleaned, fmt).replace(tzinfo=timezone.utc)
        except ValueError:
            continue
    return None


def _normal(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().lower().rstrip(".0123456789 ")


def _heading_identity(text: str) -> str:
    """Judul heading tanpa nomor, dipakai untuk mencari halaman pada TOC."""
    return re.sub(r"\s+", " ", NUMBER_RE.sub("", text)).strip().lower()


def _paragraph_pages(document: Document) -> list[tuple[int, int, str, str]]:
    """Ambil halaman render terakhir yang disimpan Word untuk setiap paragraf."""
    page = 1
    rows = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        rendered_breaks = len(paragraph._p.xpath(".//w:lastRenderedPageBreak"))
        explicit_breaks = len(paragraph._p.xpath('.//w:br[@w:type="page"]'))
        page += rendered_breaks + explicit_breaks
        rows.append((index, page, " ".join(paragraph.text.split()), paragraph.style.name if paragraph.style else ""))
    return rows


def _toc_pages(paragraphs: list[tuple[int, int, str, str]]) -> dict[str, int]:
    pages: dict[str, int] = {}
    for _, _, text, style in paragraphs:
        if not style.lower().startswith("toc"):
            continue
        match = re.match(r"^(.*?)\s+(\d+)$", text)
        if match:
            pages[_heading_identity(match.group(1))] = int(match.group(2))
    return pages


def _heading_level(style: str) -> int | None:
    match = re.fullmatch(r"heading\s*([1-6])", style.strip(), re.I)
    return int(match.group(1)) if match else None


def _heading_number(text: str) -> tuple[str, tuple[str, ...]] | None:
    match = NUMBER_RE.match(text)
    if not match:
        return None
    raw = match.group(1).rstrip(".")
    return raw, tuple(part.lower() for part in raw.split("."))


def _component_at(kind: str, position: int) -> str | None:
    if kind == "numeric":
        return str(position)
    if kind == "alpha" and 1 <= position <= 26:
        return chr(ord("a") + position - 1)
    return None


def _numbering_findings(headings: list[tuple[int, str, str]]) -> list[str]:
    findings: list[str] = []
    seen: dict[tuple[str, ...], tuple[int, str]] = {}
    latest_at_level: dict[int, tuple[str, ...]] = {}
    sibling_counts: dict[tuple[int, tuple[str, ...]], int] = {}
    sibling_kinds: dict[tuple[int, tuple[str, ...]], str] = {}

    for page, text, style in headings:
        level = _heading_level(style)
        parsed = _heading_number(text)
        if level is None:
            continue
        if parsed is None:
            findings.append(f"Halaman {page} ({style}): heading belum memiliki nomor bab: {text}")
            continue
        raw, parts = parsed
        if len(parts) != level:
            findings.append(
                f"Halaman {page}: nomor {raw} memiliki {len(parts)} tingkat, "
                f"tetapi menggunakan {style}. Gunakan Heading {len(parts)} atau perbaiki nomornya."
            )

        if parts in seen:
            first_page, first_text = seen[parts]
            findings.append(
                f"Halaman {page}: nomor {raw} duplikat; sudah dipakai pada halaman "
                f"{first_page} ({first_text})."
            )
        else:
            seen[parts] = (page, text)

        parent = parts[:-1]
        expected_parent = latest_at_level.get(level - 1) if level > 1 else ()
        if level > 1 and len(parts) == level:
            if expected_parent is None:
                findings.append(f"Halaman {page}: {raw} tidak memiliki Heading {level - 1} induk sebelumnya.")
            elif parent != expected_parent:
                expected = ".".join((*expected_parent, parts[-1]))
                findings.append(
                    f"Halaman {page}: parent nomor {raw} tidak sesuai heading induk. "
                    f"Seharusnya berada di bawah {'.'.join(expected_parent)} (contoh: {expected})."
                )

        canonical_parent = expected_parent if expected_parent is not None else parent
        sibling_key = (level, canonical_parent)
        position = sibling_counts.get(sibling_key, 0) + 1
        sibling_counts[sibling_key] = position
        kind = sibling_kinds.setdefault(
            sibling_key, "numeric" if parts[-1].isdigit() else "alpha"
        )
        expected_component = _component_at(kind, position)
        expected_parts = (*canonical_parent, expected_component) if expected_component else parts
        if expected_component is not None and parts != expected_parts:
            expected_number = ".".join(expected_parts)
            findings.append(
                f"Halaman {page}: urutan ke-{position} untuk {style} adalah {raw}; "
                f"seharusnya {expected_number}."
            )
        latest_at_level[level] = tuple(expected_parts)
        for deeper in range(level + 1, 7):
            latest_at_level.pop(deeper, None)

    return findings


def _section_page(heading: str, page_map: dict[str, int], fallback_map: dict[str, int]) -> int:
    return page_map.get(_heading_identity(heading), fallback_map.get(heading, 1))


def _diagram_content_findings(
    sections: list[dict[str, Any]],
    toc_page_map: dict[str, int],
    heading_page_map: dict[str, int],
) -> list[str]:
    findings: list[str] = []
    for section in sections:
        heading = section.get("heading") or ""
        parsed_heading = docx_parser.parse_diagram_heading(heading)
        if not parsed_heading:
            continue
        kind, sp_name, is_summary = parsed_heading
        if is_summary or not sp_name:
            continue
        page = _section_page(heading, toc_page_map, heading_page_map)
        items = section.get("items", [])
        image_positions = [index for index, (kind_, value) in enumerate(items) if kind_ == "image" and docx_parser._is_diagram(value)]
        label = "Data Model" if kind == "data_model" else "Data Flow"
        if not image_positions:
            findings.append(f"Halaman {page}: {sp_name} tidak memiliki diagram {label} yang valid.")
            continue

        first_image = image_positions[0]
        marker_position = next((
            index for index, (kind_, value) in enumerate(items)
            if index > first_image and kind_ == "text" and C.EXPLANATION_HINT_RE.search(value)
        ), None)
        if marker_position is None:
            findings.append(
                f"Halaman {page}: {label} {sp_name} memiliki gambar, tetapi subbagian "
                "interpretasi atau penjelasan setelah diagram tidak ditemukan."
            )
            continue
        explanation = " ".join(
            value for kind_, value in items[marker_position + 1:] if kind_ == "text"
        ).strip()
        if len(explanation) < 80:
            findings.append(
                f"Halaman {page}: interpretasi {label} {sp_name} terlalu singkat atau kosong "
                "setelah judul penjelasan."
            )
    return findings


def _spec_pair_findings(
    sections: list[dict[str, Any]],
    toc_page_map: dict[str, int],
    heading_page_map: dict[str, int],
) -> list[str]:
    findings: list[str] = []
    for section in sections:
        heading = section.get("heading") or ""
        low = heading.lower()
        category = next((label for key, label in (("table source", "Table Source"), ("table parameter", "Table Parameter"), ("table target", "Table Target")) if key in low), None)
        if category is None:
            continue
        page = _section_page(heading, toc_page_map, heading_page_map)
        pending_name: str | None = None
        valid_pairs = 0
        for kind, value in section.get("items", []):
            if kind == "text":
                match = C.SPEC_TABLE_LABEL_RE.search(value)
                if not match:
                    continue
                if pending_name is not None:
                    findings.append(
                        f"Halaman {page}: {category} {pending_name} hanya memiliki judul; "
                        "tabel spesifikasi Field Name, Data Type, Description tidak ditemukan."
                    )
                pending_name = match.group("name")
            elif kind == "table" and value:
                header = [cell.strip().lower() for cell in value[0]]
                if header[:len(C.SPEC_TABLE_HEADER)] != C.SPEC_TABLE_HEADER:
                    continue
                if pending_name is None:
                    findings.append(f"Halaman {page}: ditemukan tabel spesifikasi tanpa nama tabel pada {category}.")
                    continue
                valid_pairs += 1
                pending_name = None
        if pending_name is not None:
            findings.append(
                f"Halaman {page}: {category} {pending_name} hanya memiliki judul; "
                "tabel spesifikasi Field Name, Data Type, Description tidak ditemukan."
            )
        if valid_pairs == 0:
            findings.append(f"Halaman {page}: {category} belum memiliki pasangan judul dan tabel spesifikasi yang valid.")
    return findings


def check_document(temp_path: Path, filename: str) -> dict[str, Any]:
    name = validate_filename(filename)
    validate_docx(temp_path)
    document = Document(temp_path)
    segment = docx_parser.segment_from_filename(Path(filename).stem)
    parsed = docx_parser.parse(temp_path, PREVIEW_IMAGES_DIR, segment_override=segment)
    sections = docx_parser._sections(document)
    paragraphs = _paragraph_pages(document)
    text_paragraphs = [(index, page, text, style) for index, page, text, style in paragraphs if text]
    toc_page_map = _toc_pages(paragraphs)
    headings = [
        (toc_page_map.get(_heading_identity(text), page), text, style)
        for _, page, text, style in text_paragraphs
        if style.lower().replace(" ", "").startswith("heading")
    ]
    heading_page_map = {
        text: page for _, page, text, style in text_paragraphs
        if style.lower().replace(" ", "").startswith("heading")
    }
    checks: list[dict[str, Any]] = []

    checks.append(_result("integrity", "Integritas DOCX", True, "Paket DOCX dapat dibuka dan bagian utama tersedia."))
    filename_ok = bool(re.match(r"^NTT[_ ]*Data[_ ]*Draft[_ ]*(?:TSD[_ ]*)?.+\.docx$", name, re.I))
    checks.append(_result("filename", "Nama dan identitas dokumen", filename_ok, "Nama file mengikuti pola sumber TSD." if filename_ok else "Nama file belum mengikuti pola NTT_Data_Draft_[segment].docx.", [] if filename_ok else [f"Nama saat ini: {name}"]))

    now = datetime.now(timezone.utc)
    created = document.core_properties.created
    modified = document.core_properties.modified
    if created and created.tzinfo is None: created = created.replace(tzinfo=timezone.utc)
    if modified and modified.tzinfo is None: modified = modified.replace(tzinfo=timezone.utc)
    metadata_ok = bool(created and modified and created <= modified <= now)
    metadata_findings = []
    if not created: metadata_findings.append("Metadata tanggal dibuat belum tersedia.")
    if not modified: metadata_findings.append("Metadata tanggal perubahan belum tersedia.")
    if created and modified and created > modified: metadata_findings.append("Tanggal dibuat lebih baru daripada tanggal perubahan.")
    if modified and modified > now: metadata_findings.append("Tanggal perubahan berada di masa depan.")
    checks.append(_result("metadata", "Metadata dan tanggal file", metadata_ok, "Tanggal dibuat dan diperbarui konsisten." if metadata_ok else "Metadata tanggal perlu diperbaiki.", metadata_findings))

    change_rows = parsed.get("change_control", [])
    change_findings: list[str] = []
    dates = []
    for index, row in enumerate(change_rows, 2):
        if len(row) < 4 or any(not str(cell).strip() for cell in row[:4]):
            change_findings.append(f"Change control baris {index} belum mengisi Version, Date, Authors, dan Summary of Changes.")
            continue
        parsed_date = _parse_date(str(row[1]))
        if parsed_date is None:
            change_findings.append(f"Tanggal change control tidak dikenali pada baris {index}: {row[1]}")
        else:
            dates.append(parsed_date)
            if parsed_date > now: change_findings.append(f"Tanggal change control berada di masa depan pada baris {index}.")
    if not change_rows: change_findings.append("Tabel Document Change Control tidak ditemukan.")
    if dates and dates != sorted(dates): change_findings.append("Urutan tanggal change control belum kronologis.")
    checks.append(_result("change_control", "Document change control", not change_findings, f"{len(change_rows)} riwayat perubahan terbaca." if change_rows else "Riwayat perubahan tidak tersedia.", change_findings))

    toc_rows = [(index, text) for index, _, text, style in text_paragraphs if style.lower().startswith("toc") and "table of contents" not in text.lower()]
    toc_names = {_normal(re.sub(r"\s+\d+$", "", text)) for _, text in toc_rows}
    missing_toc = [text for _, text, _ in headings if _normal(text) not in toc_names]
    toc_findings = [f"Heading belum tercantum di TOC: {text}" for text in missing_toc[:12]]
    if not toc_rows: toc_findings.insert(0, "Daftar isi dengan style TOC tidak ditemukan atau belum diperbarui.")
    checks.append(_result("toc", "Table of contents", bool(toc_rows) and not missing_toc, f"{len(toc_rows)} entri TOC cocok dengan heading." if not missing_toc else f"{len(missing_toc)} heading belum sinkron dengan TOC.", toc_findings))

    required = ("system architecture", "data design", "appendix")
    top = [(page, text.lower()) for page, text, style in headings if style.lower().replace(" ", "") == "heading1"]
    positions = [next((i for i, text in top if title in text), None) for title in required]
    chapter_ok = all(pos is not None for pos in positions) and positions == sorted(positions)
    chapter_findings = [f"Bab wajib tidak ditemukan: {title.title()}" for title, pos in zip(required, positions) if pos is None]
    if all(pos is not None for pos in positions) and positions != sorted(positions): chapter_findings.append("Urutan bab wajib harus System Architecture, Data Design, lalu Appendix.")
    checks.append(_result("chapters", "Bab wajib dan urutan", chapter_ok, "Tiga bab utama tersedia dalam urutan yang benar." if chapter_ok else "Struktur bab utama belum sesuai.", chapter_findings))

    hierarchy_findings = []
    previous_level = 0
    for page, text, style in headings:
        match = re.search(r"(\d+)$", style)
        level = int(match.group(1)) if match else 0
        if previous_level and level > previous_level + 1:
            hierarchy_findings.append(f"Halaman {page}: level heading meloncat dari {previous_level} ke {level} ({text}).")
        previous_level = level
    checks.append(_result("hierarchy", "Hierarki heading", bool(headings) and not hierarchy_findings, f"{len(headings)} heading memiliki hierarki yang konsisten." if headings else "Heading tidak ditemukan.", hierarchy_findings or ([] if headings else ["Gunakan style Heading 1, Heading 2, dan seterusnya."])))

    number_findings = _numbering_findings(headings)
    checks.append(_result("numbering", "Penomoran Heading 1–6", bool(headings) and not number_findings, f"Urutan dan parent number {len(headings)} heading konsisten hingga Heading 6." if not number_findings else f"Ditemukan {len(number_findings)} masalah urutan, duplikasi, atau parent number.", number_findings[:40]))

    procedures = parsed.get("procedures", [])
    procedure_findings = _diagram_content_findings(sections, toc_page_map, heading_page_map)
    for procedure in procedures:
        for kind, label in (("data_model", "Data Model"), ("data_flow", "Data Flow")):
            block = procedure.get(kind) or {}
            if not block.get("images"): procedure_findings.append(f"{procedure['sp_name']}: diagram {label} tidak ditemukan.")
    if not procedures: procedure_findings.append("Tidak ada stored procedure yang dapat dikenali dari heading Data Model/Data Flow.")
    checks.append(_result("procedures", "Kelengkapan stored procedure", bool(procedures) and not procedure_findings, f"{len(procedures)} SP memiliki pasangan diagram dan penjelasan." if not procedure_findings else f"{len(procedure_findings)} bagian SP belum lengkap.", procedure_findings[:30]))

    specs = parsed.get("table_specs", {})
    spec_findings = _spec_pair_findings(sections, toc_page_map, heading_page_map)
    for key, label in (("source", "Table Source"), ("parameter", "Table Parameter"), ("target", "Table Target")):
        tables = (specs.get(key) or {}).get("tables", [])
        if not tables and not any(label in finding for finding in spec_findings):
            spec_findings.append(f"Appendix {label} belum memiliki tabel spesifikasi.")
        for table in tables:
            if not table.get("table_name"): spec_findings.append(f"{label}: nama tabel belum diisi.")
            if not table.get("fields"): spec_findings.append(f"{label} {table.get('table_name') or '(tanpa nama)'}: field belum tersedia.")
            for row, field in enumerate(table.get("fields", []), 2):
                if not all(field.get(key) for key in ("field", "type", "description")):
                    spec_findings.append(f"{label} {table.get('table_name')}: baris {row} belum lengkap.")
    checks.append(_result("table_specs", "Spesifikasi tabel", not spec_findings, "Table Source, Parameter, dan Target memiliki field lengkap." if not spec_findings else "Spesifikasi tabel masih memiliki bagian kosong.", spec_findings[:30]))

    placeholder_findings = [f"Halaman {page}: teks sementara '{match.group(0)}'." for _, page, text, _ in text_paragraphs for match in [PLACEHOLDER_RE.search(text)] if match]
    checks.append(_result("placeholders", "Teks sementara dan kelengkapan isi", not placeholder_findings, "Tidak ditemukan placeholder penyuntingan." if not placeholder_findings else "Masih ada teks sementara yang harus diselesaikan.", placeholder_findings[:20]))

    parser_findings = list(parsed.get("warnings", []))
    checks.append(_result("parser", "Keterbacaan oleh mesin", not parser_findings, "Dokumen dapat dipindai tanpa peringatan parser." if not parser_findings else f"Parser menemukan {len(parser_findings)} masalah.", parser_findings[:20]))

    passed = sum(check["status"] == "pass" for check in checks)
    score = round(passed / len(checks) * 100)
    report_id = uuid.uuid4().hex
    STAGING_DIR.mkdir(parents=True, exist_ok=True)
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    staged_path = STAGING_DIR / f"{report_id}.docx"
    # Upload berada di /tmp, sedangkan staging berada di volume Docker.
    # shutil.move menangani batas filesystem yang tidak didukung os.replace.
    shutil.move(str(temp_path), staged_path)
    report = {
        "id": report_id, "checker_version": CHECKER_VERSION, "filename": name, "segment": parsed.get("segment"),
        "score": score, "eligible": score == 100, "status": "ready" if score == 100 else "needs_revision",
        "checked_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "summary": {"checks": len(checks), "passed": passed, "failed": len(checks) - passed, "procedures": len(procedures)},
        "checks": checks,
    }
    (REPORT_DIR / f"{report_id}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report


def list_reports(limit: int | None = None) -> list[dict[str, Any]]:
    REPORT_DIR.mkdir(parents=True, exist_ok=True)
    reports = []
    for path in REPORT_DIR.glob("*.json"):
        try: reports.append(json.loads(path.read_text(encoding="utf-8")))
        except (OSError, json.JSONDecodeError): continue
    ordered = sorted(reports, key=lambda item: item.get("checked_at", ""), reverse=True)
    return ordered[:limit] if limit is not None else ordered


def load_report(report_id: str) -> dict[str, Any] | None:
    if not re.fullmatch(r"[a-f0-9]{32}", report_id): return None
    path = REPORT_DIR / f"{report_id}.json"
    if not path.exists(): return None
    return json.loads(path.read_text(encoding="utf-8"))


def mark_activated(report: dict[str, Any]) -> None:
    report["status"] = "active"
    report["activated_at"] = datetime.now(timezone.utc).isoformat(timespec="seconds")
    (REPORT_DIR / f"{report['id']}.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")


def recheck_report(report_id: str) -> dict[str, Any]:
    report = load_report(report_id)
    if report is None:
        raise FileNotFoundError("Hasil pemeriksaan tidak ditemukan.")
    if report.get("status") == "active":
        raise ValueError("Dokumen aktif tidak memiliki file staging untuk diperiksa ulang.")
    staged = STAGING_DIR / f"{report_id}.docx"
    if not staged.exists():
        raise FileNotFoundError("File staging sudah tidak tersedia.")
    refreshed = check_document(staged, report["filename"])
    (REPORT_DIR / f"{report_id}.json").unlink(missing_ok=True)
    return refreshed
